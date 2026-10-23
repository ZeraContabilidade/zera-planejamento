"""
Relatório em Excel (.xlsx) com gráficos nativos (editáveis no Excel) e memória de cálculo.

Abas: Resumo · Histórico mensal · uma aba por regime · Base mensal · Origem dos dados ·
Projeções · Transição 2026-2033 · Equilíbrio · Sócios e empresa · Prazos · Premissas · Alertas.
"""
from __future__ import annotations

from dataclasses import asdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..base import COLUNAS
from ..calculo import COMPONENTES, NOME_TRIBUTO, ROTULO
from ..configuracoes import Configuracoes
from ..formatos import brl, comp, pct
from ..projecao import CENARIOS
from .graficos import COR

PRETO, DOURADO, DOURADO_TXT, CLARO, CINZA = "111111", "C5A15F", "A8843F", "F6F0E4", "5F5F5F"
FONTE = "Inter"
RS = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
PCT = "0.00%"
FINO = Side(style="thin", color="DDDDDD")


def _cab(ws, linha: int, valores: list[str], col0: int = 1) -> None:
    for j, v in enumerate(valores, start=col0):
        c = ws.cell(row=linha, column=j, value=v)
        c.fill = PatternFill("solid", fgColor=PRETO)
        c.font = Font(name=FONTE, bold=True, color=DOURADO, size=9)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[linha].height = 30


def _celula(ws, linha, col, valor, fmt=None, negrito=False, fundo=None):
    c = ws.cell(row=linha, column=col, value=valor)
    c.font = Font(name=FONTE, size=9, bold=negrito)
    c.border = Border(bottom=FINO)
    if fmt:
        c.number_format = fmt
    if fundo:
        c.fill = PatternFill("solid", fgColor=fundo)
    return c


def _larguras(ws, larguras: dict[str, float]) -> None:
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w


def _tabela(ws, df: pd.DataFrame, linha: int, formatos: dict[str, str] | None = None, indice: str | None = None) -> int:
    """Escreve um DataFrame com cabeçalho Zera; devolve a próxima linha livre."""
    cols = ([indice] if indice else []) + list(df.columns)
    _cab(ws, linha, cols)
    for i, (idx, row) in enumerate(df.iterrows(), start=linha + 1):
        j = 1
        if indice:
            _celula(ws, i, j, comp(idx) if isinstance(idx, pd.Period) else idx)
            j += 1
        for col, v in row.items():
            v = None if isinstance(v, float) and pd.isna(v) else v
            _celula(ws, i, j, v, (formatos or {}).get(col, RS if isinstance(v, (int, float)) and not isinstance(v, bool) else None))
            j += 1
    ws.freeze_panes = ws.cell(row=linha + 1, column=2)
    return linha + len(df) + 2


def exportar_excel(pl, caminho: Path, graficos_png: dict[str, Path] | None = None) -> Path:
    cfg = Configuracoes.carregar()
    c = pl.cliente
    wb = Workbook()

    # ---------------- Resumo ----------------
    ws = wb.active
    ws.title = "Resumo"
    ws.sheet_view.showGridLines = False
    if cfg.kit and (cfg.kit / "logo_zera_fundo_preto.png").exists():
        img = XLImage(str(cfg.kit / "logo_zera_fundo_preto.png"))
        img.height, img.width = 48, int(48 * img.width / img.height)
        ws.add_image(img, "H1")
    ws["A1"] = f"Planejamento tributário — {c.nome}"
    ws["A1"].font = Font(name="Cormorant Garamond", size=18, bold=True, color=PRETO)
    mh, mp = pl.meses_historico, pl.meses_projetados
    ws["A2"] = (f"CNPJ {c.cnpj or '–'} · {c.municipio}/{c.uf} · regime atual: {c.regime_atual} · histórico "
                f"{comp(mh[0])} a {comp(mh[-1])} · projeção {comp(mp[0])} a {comp(mp[-1])}")
    ws["A2"].font = Font(name=FONTE, size=9, color=CINZA)
    m = pl.melhor("Neutro")
    if m is not None:
        ws["A3"] = (f"Cenário Neutro: menor carga no {m['rotulo']} — {brl(m['tributos'])} ({pct(m['carga'])} da receita), "
                    f"{brl(m['economia_vs_2o'])} abaixo do 2º colocado.")
        ws["A3"].font = Font(name=FONTE, size=10, bold=True, color=DOURADO_TXT)
    ws["A4"] = f"Rascunho técnico gerado em {pl.gerado_em:%d/%m/%Y %H:%M} — revisar antes de enviar ao cliente."
    ws["A4"].font = Font(name=FONTE, size=8, italic=True, color=CINZA)

    r0 = 6
    _cab(ws, r0, ["Cenário", "Regime", "Receita", "Tributos", "Carga efetiva", "Situação", "Economia vs. 2º"])
    for i, lin in enumerate(pl.resumo.itertuples(), start=r0 + 1):
        destaque = CLARO if lin.melhor else None
        _celula(ws, i, 1, lin.cenario, negrito=lin.melhor, fundo=destaque)
        _celula(ws, i, 2, lin.rotulo, negrito=lin.melhor, fundo=destaque)
        _celula(ws, i, 3, lin.receita, RS, fundo=destaque)
        _celula(ws, i, 4, None if pd.isna(lin.tributos) else lin.tributos, RS, lin.melhor, destaque)
        _celula(ws, i, 5, None if pd.isna(lin.carga) else lin.carga, PCT, lin.melhor, destaque)
        _celula(ws, i, 6, "MENOR CARGA" if lin.melhor else ("" if lin.elegivel else "não aplicável"), negrito=lin.melhor,
                fundo=destaque)
        _celula(ws, i, 7, lin.economia_vs_2o if lin.melhor else None, RS, fundo=destaque)
    fim_tab = r0 + len(pl.resumo)

    # tabela auxiliar para o gráfico de barras (cenário × regime)
    piv = pl.resumo.pivot(index="cenario", columns="regime", values="tributos").reindex(list(CENARIOS))
    aux = fim_tab + 3
    ws.cell(row=aux - 1, column=1, value="Tributos por cenário (base do gráfico)").font = Font(name=FONTE, bold=True, size=9)
    _cab(ws, aux, ["Cenário"] + [ROTULO[r] for r in piv.columns])
    for i, (cen, row) in enumerate(piv.iterrows(), start=aux + 1):
        _celula(ws, i, 1, cen)
        for j, v in enumerate(row, start=2):
            _celula(ws, i, j, None if pd.isna(v) else v, RS)
    bar = BarChart()
    bar.type, bar.grouping, bar.title = "col", "clustered", "Custo tributário do período projetado"
    bar.y_axis.numFmt = '"R$" #,##0'
    bar.add_data(Reference(ws, min_col=2, max_col=1 + len(piv.columns), min_row=aux, max_row=aux + 3), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=aux + 1, max_row=aux + 3))
    for s, r in zip(bar.series, piv.columns):
        s.graphicalProperties.solidFill = COR[r].lstrip("#")
        s.graphicalProperties.line.noFill = True
    bar.height, bar.width = 8, 18
    ws.add_chart(bar, f"I{r0}")

    eq = aux + 6
    ws.cell(row=eq - 1, column=1, value="Pontos de equilíbrio (cenário Neutro)").font = Font(name=FONTE, bold=True, size=10,
                                                                                         color=DOURADO_TXT)
    _cab(ws, eq, ["Comparação", "Indicador", "Equilíbrio", "Projetado", "Leitura"])
    for i, e in enumerate(pl.equilibrio.itertuples(), start=eq + 1):
        fmt = PCT if e.unidade == "%" else RS
        _celula(ws, i, 1, e.comparacao)
        _celula(ws, i, 2, e.indicador)
        _celula(ws, i, 3, None if pd.isna(e.equilibrio) else e.equilibrio, fmt)
        _celula(ws, i, 4, None if pd.isna(e.projetado) else e.projetado, fmt)
        cel = _celula(ws, i, 5, e.leitura)
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=5, end_row=i, end_column=8)
        ws.row_dimensions[i].height = 42
    _larguras(ws, {"A": 14, "B": 30, "C": 16, "D": 16, "E": 13, "F": 14, "G": 16, "H": 12})

    # ---------------- Histórico mensal ----------------
    wh = wb.create_sheet("Histórico mensal")
    mensal = pl.historico.mensal().loc[pl.meses_historico]
    mensal.insert(0, "Receita", pl.base.receita_bruta().loc[pl.meses_historico])
    mensal.columns = ["Receita"] + [ROTULO[r] for r in mensal.columns[1:]]
    _tabela(wh, mensal, 1, indice="Competência")
    linha = LineChart()
    linha.title, linha.y_axis.numFmt = "Tributos mês a mês por regime", '"R$" #,##0'
    linha.add_data(Reference(wh, min_col=3, max_col=1 + mensal.shape[1], min_row=1, max_row=1 + len(mensal)),
                   titles_from_data=True)
    linha.set_categories(Reference(wh, min_col=1, min_row=2, max_row=1 + len(mensal)))
    for s, r in zip(linha.series, pl.historico.mensal().columns):
        s.graphicalProperties.line.solidFill = COR[r].lstrip("#")
        s.graphicalProperties.line.width = 22000
        s.smooth = False
    linha.height, linha.width = 9, 22
    wh.add_chart(linha, f"{get_column_letter(mensal.shape[1] + 3)}2")
    _larguras(wh, {get_column_letter(i): 18 for i in range(1, mensal.shape[1] + 2)})

    # ---------------- Uma aba por regime ----------------
    for reg in pl.historico.regimes:
        wr = wb.create_sheet(ROTULO[reg][:31])
        blocos = [pl.historico.regimes[reg].loc[pl.meses_historico].assign(periodo="Histórico")]
        for nome in CENARIOS:
            res = pl.cenarios[nome].resultado.regimes.get(reg)
            if res is not None:
                blocos.append(res.loc[pl.cenarios[nome].meses].assign(periodo=f"Projeção {nome}"))
        df = pd.concat(blocos)
        cols = ["periodo", "receita"] + [k for k in COMPONENTES[reg] if k in df] + ["total"]
        extra = [k for k in df.columns if k not in cols and k not in ("elegivel", "hibrido_ok")]
        df = df[cols + extra].rename(columns={**NOME_TRIBUTO, "periodo": "Período", "receita": "Receita",
                                              "total": "TOTAL"})
        fmts = {k: PCT for k in df.columns if str(k).startswith("aliq") or k == "fator_r"}
        fmts.update({k: "@" for k in ("anexo_fator_r",)})
        fmts["faixa"] = "0"
        _tabela(wr, df, 1, fmts, indice="Competência")
        _larguras(wr, {get_column_letter(i): 15 for i in range(1, df.shape[1] + 2)})

    # ---------------- Base e origem ----------------
    wbm = wb.create_sheet("Base mensal")
    base = pl.base.dados.copy()
    base.columns = [COLUNAS[col][1] for col in base.columns]
    _tabela(wbm, base, 1, indice="Competência")
    _larguras(wbm, {get_column_letter(i): 18 for i in range(1, base.shape[1] + 2)})
    wo = wb.create_sheet("Origem dos dados")
    org = pl.base.origem.copy()
    org.columns = [COLUNAS[col][1] for col in org.columns]
    _tabela(wo, org, 1, indice="Competência")

    # ---------------- Projeções ----------------
    wp = wb.create_sheet("Projeções")
    proj = pd.concat([pl.cenarios[n].base_projetada.dados.assign(cenario=n) for n in CENARIOS])
    proj = proj[["cenario"] + list(COLUNAS)].dropna(axis=1, how="all")
    proj.columns = ["Cenário"] + [COLUNAS[col][1] for col in proj.columns[1:]]
    _tabela(wp, proj, 1, indice="Competência")

    # ---------------- Transição ----------------
    if pl.transicao is not None and not pl.transicao.empty:
        wt = wb.create_sheet("Transição 2026-2033")
        t = pl.transicao.copy()
        t.columns = [ROTULO.get(col, "Receita") for col in t.columns]
        _tabela(wt, t, 1, indice="Ano")
        ch = LineChart()
        ch.title, ch.y_axis.numFmt = "Carga anual por regime na transição", '"R$" #,##0'
        n_reg = len([col for col in pl.transicao.columns if col in ROTULO])
        ch.add_data(Reference(wt, min_col=2, max_col=1 + n_reg, min_row=1, max_row=1 + len(t)), titles_from_data=True)
        ch.set_categories(Reference(wt, min_col=1, min_row=2, max_row=1 + len(t)))
        for s, r in zip(ch.series, [col for col in pl.transicao.columns if col in ROTULO]):
            s.graphicalProperties.line.solidFill = COR[r].lstrip("#")
        ch.height, ch.width = 9, 20
        wt.add_chart(ch, "H2")
        _larguras(wt, {get_column_letter(i): 18 for i in range(1, 7)})

    # ---------------- Equilíbrio (curva) ----------------
    if pl.curva is not None and not pl.curva.empty:
        we = wb.create_sheet("Equilíbrio")
        cv = pl.curva.set_index("margem")
        cv.columns = [ROTULO[col] for col in cv.columns]
        _tabela(we, cv, 1, indice="Margem operacional")
        for i in range(2, len(cv) + 2):
            we.cell(row=i, column=1).number_format = PCT
        _larguras(we, {get_column_letter(i): 20 for i in range(1, 6)})

    # ---------------- Sócios + empresa (Lei 15.270/2025) ----------------
    if pl.socios is not None and not pl.socios.vazia:
        wsoc = wb.create_sheet("Sócios e empresa")
        r = pl.socios.resumo[["rotulo", "tributos_pj", "lucro_contabil", "aliquota_pj", "dividendos", "irrf_dividendos",
                              "irpfm", "inss_socios", "irrf_pro_labore", "custo_pf", "carga_total", "carga_pct"]].copy()
        r["melhor"] = pl.socios.resumo["melhor"].map({True: "MENOR CARGA", False: ""})
        r.columns = ["Regime", "Tributos da empresa", "Lucro contábil", "Alíquota IRPJ+CSLL", "Lucros distribuídos",
                     "IRRF 10% (antecipação)", "IRPFM", "INSS dos sócios", "IRRF do pró-labore", "Custo dos sócios",
                     "Empresa + sócios", "Carga total", "Situação"]
        prox = _tabela(wsoc, r.set_index("Regime"), 1, {"Alíquota IRPJ+CSLL": PCT, "Carga total": PCT}, indice="Regime")
        det = pl.socios.detalhe[["rotulo", "socio", "participacao", "pro_labore_mes", "inss", "irrf_pro_labore",
                                 "dividendos", "dividendos_mes", "irrf_dividendos", "renda_anual", "aliquota_irpfm",
                                 "redutor", "irpfm", "ajuste_anual"]].copy()
        det.columns = ["Regime", "Sócio", "Participação", "Pró-labore/mês", "INSS", "IRRF pró-labore", "Lucros",
                       "Lucros/mês", "IRRF 10%", "Renda anual", "Alíquota IRPFM", "Redutor", "IRPFM", "Ajuste anual"]
        _tabela(wsoc, det.set_index("Regime"), prox + 1, {"Participação": PCT, "Alíquota IRPFM": PCT}, indice="Regime")
        linha = prox + len(det) + 3
        for o in pl.socios.observacoes:
            _celula(wsoc, linha, 1, o)
            linha += 1
        _larguras(wsoc, {get_column_letter(i): 18 for i in range(1, 15)})

    # ---------------- Prazos ----------------
    if pl.prazos is not None and not pl.prazos.empty:
        wz = wb.create_sheet("Prazos")
        pz = pl.prazos.copy()
        pz["data"] = pd.to_datetime(pz["data"]).dt.strftime("%d/%m/%Y")
        pz.columns = ["Data", "Dias", "Evento", "Base legal", "Ação"]
        _tabela(wz, pz.set_index("Data"), 1, {"Dias": "0"}, indice="Data")
        _larguras(wz, {"A": 12, "B": 8, "C": 80, "D": 50, "E": 14})

    # ---------------- Premissas e alertas ----------------
    wpr = wb.create_sheet("Premissas")
    prem = [(k, str(v)) for k, v in asdict(c).items() if k not in ("depara", "socios")]
    prem += [("socio", f"{s.nome} — {pct(s.participacao)}") for s in c.socios]
    prem += [("regras_pendentes_conferir", p) for p in pl.pendencias]
    dfp = pd.DataFrame(prem, columns=["Parâmetro", "Valor"]).set_index("Parâmetro")
    _tabela(wpr, dfp, 1, indice="Parâmetro")
    _larguras(wpr, {"A": 32, "B": 120})
    wa = wb.create_sheet("Alertas")
    al = pl.alertas.como_df()
    if al.empty:
        al = pd.DataFrame({"mensagem": ["Sem alertas."]})
    al.columns = [col.capitalize() for col in al.columns]
    _tabela(wa, al, 1)
    _larguras(wa, {"A": 12, "B": 14, "C": 140})

    for w in wb.worksheets:
        w.sheet_properties.tabColor = PRETO if w.title == "Resumo" else DOURADO
    wb.save(caminho)
    return caminho
