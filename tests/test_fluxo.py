"""Base, projeção, equilíbrio, relatórios, linha de comando e app — fluxo completo com dados sintéticos."""
import math

import openpyxl
import pandas as pd
import pytest

from zera_planejamento.base import BaseMensal
from zera_planejamento.planejamento import gerar_planejamento
from zera_planejamento.projecao import projetar


def test_base_salvar_carregar(tmp_path, base_exemplo):
    base_exemplo.dados.loc["2026-01", "cred_aluguel"] = 1234.56
    base_exemplo.salvar(tmp_path)
    b = BaseMensal.carregar(tmp_path)
    assert b.dados.loc["2026-01", "cred_aluguel"] == 1234.56
    assert b.origem.loc["2025-05", "rec_anexo_iii"] == "exemplo"
    assert len(b.meses) == len(base_exemplo.meses)


def test_planilha_de_digitacao(tmp_path, base_exemplo):
    arq = base_exemplo.para_planilha(tmp_path / "base.xlsx")
    df = BaseMensal.de_planilha(arq)
    assert math.isclose(df.loc["2026-09", "rec_anexo_iii"], base_exemplo.dados.loc["2026-09", "rec_anexo_iii"])


def test_projecao_cenarios(base_exemplo, cliente):
    cen = projetar(base_exemplo, cliente)
    n, o, p = cen["Neutro"], cen["Otimista"], cen["Pessimista"]
    assert list(n.meses.astype(str)) == ["2026-10", "2026-11", "2026-12", "2027-01", "2027-02", "2027-03"]
    rn = n.base_projetada.receita_bruta().sum()
    assert math.isclose(o.base_projetada.receita_bruta().sum(), rn * 1.15)
    assert math.isclose(p.base_projetada.receita_bruta().sum(), rn * 0.85)
    # tendência linear de +2.000/mês continua: out/26 = 80.000 + 2.000 × 21
    assert math.isclose(n.base_projetada.dados.loc["2026-10", "rec_anexo_iii"], 122_000, rel_tol=1e-9)
    # folha é fixa (média dos últimos 3 meses) e não muda com o cenário
    assert (o.base_projetada.dados["salarios"] == 15_000).all()


def test_planejamento_completo(base_exemplo, cliente):
    pl = gerar_planejamento(cliente, base_exemplo)
    assert set(pl.resumo.cenario) == {"Neutro", "Otimista", "Pessimista"}
    assert pl.resumo.groupby("cenario").melhor.sum().eq(1).all()
    assert {"Lucro Real × Lucro Presumido", "Simples × Lucro Presumido"} <= set(pl.equilibrio.comparacao)
    assert list(pl.transicao.index) == list(range(2026, 2034))
    assert not pl.curva.empty


def test_relatorios(tmp_path, base_exemplo, cliente, monkeypatch):
    monkeypatch.setenv("ZERA_DADOS", str(tmp_path / "dados"))
    from zera_planejamento.cliente import Socio
    from zera_planejamento.relatorio import gerar_relatorios
    cliente.socios = [Socio("Sócio Um", 0.7), Socio("Sócio Dois", 0.3)]
    pl = gerar_planejamento(cliente, base_exemplo)
    arqs = gerar_relatorios(pl, tmp_path / "rel")
    assert arqs["pdf"].stat().st_size > 50_000
    wb = openpyxl.load_workbook(arqs["excel"])
    assert {"Resumo", "Histórico mensal", "Simples Nacional", "Lucro Presumido", "Lucro Real", "Base mensal",
            "Transição 2026-2033", "Sócios e empresa", "Prazos", "Premissas", "Alertas"} <= set(wb.sheetnames)
    assert wb["Resumo"]._charts and wb["Histórico mensal"]._charts


def test_cli(tmp_path, monkeypatch, relatorios, cliente, capsys):
    monkeypatch.setenv("ZERA_DADOS", str(tmp_path / "dados"))
    from zera_planejamento.cli import main
    cliente.salvar()
    assert main(["importar", "--cliente", "11222333", *map(str, relatorios.values())]) == 0
    assert main(["clientes"]) == 0
    assert "Empresa Teste" in capsys.readouterr().out


def test_app_todas_as_paginas(tmp_path, monkeypatch, base_exemplo, cliente):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setenv("ZERA_DADOS", str(tmp_path / "dados"))
    from zera_planejamento.cliente import Socio
    cliente.socios = [Socio("Sócio Um", 1.0)]
    cliente.salvar()
    base_exemplo.salvar(cliente.pasta())
    import zera_planejamento.app as app_pkg
    from pathlib import Path
    main = Path(app_pkg.__file__).parent / "main.py"
    for pagina in ("inicio", "dashboard", "cadastro", "importar", "documentos", "base", "planejamento", "anos", "reforma", "simulador", "socios", "relatorios", "carteira", "depara",
                   "configuracoes"):
        at = AppTest.from_file(str(main), default_timeout=300)
        at.session_state["cliente_id"] = cliente.id
        at.run()
        at.switch_page(f"paginas/{pagina}.py").run()
        assert not at.exception, f"{pagina}: {[e.message for e in at.exception]}"


def test_painel_do_dashboard(base_exemplo, cliente):
    from zera_planejamento.painel import montar_painel
    pl = gerar_planejamento(cliente, base_exemplo)
    p = montar_painel(pl, "Otimista", "projecao")
    ind = p.indicadores
    # o recomendado é o 1º do ranking e bate com o resumo do cenário Otimista
    melhor = pl.melhor("Otimista")
    assert ind["recomendado"] == melhor["regime"] == p.ranking.iloc[0]["regime"]
    assert math.isclose(ind["tributos_recomendado"], melhor["tributos"])
    # economia anual = (regime atual − menor carga) × 12 ÷ meses projetados
    atual = p.ranking.set_index("regime").loc[cliente.regime_atual, "tributos"]
    assert math.isclose(ind["economia_ano"], (atual - melhor["tributos"]) * 12 / len(pl.meses_projetados))
    # composição soma os tributos de cada regime
    soma = p.composicao.groupby("regime")["valor"].sum()
    for r in soma.index:
        assert math.isclose(soma[r], p.ranking.set_index("regime").at[r, "tributos"], rel_tol=1e-6)
    assert set(p.mensal["fase"]) == {"Histórico", "Projeção"}
    h = montar_painel(pl, "Neutro", "historico")
    assert len(h.meses) == 12 and h.indicadores["n_meses"] == 12


def test_parecer_ao_cliente(tmp_path, base_exemplo, cliente):
    import zipfile

    import docx

    from zera_planejamento.relatorio.parecer import gerar_parecer, montar_parecer
    pl = gerar_planejamento(cliente, base_exemplo, com_transicao=False, com_equilibrio=False)
    par = montar_parecer(pl)
    melhor = pl.melhor("Neutro")
    # a linha destacada da tabela é o regime recomendado, com o custo anualizado do motor
    assert par.tabela[par.destaque][3].endswith("ecomendado")
    assert par.tabela[par.destaque][1].replace(".", "").endswith(f"{round(melhor['tributos'] * 2):,}".replace(",", ""))
    assert [n for n, _, _ in par.cards] == ["01", "02", "03"]
    arqs = gerar_parecer(pl, tmp_path)
    texto = "\n".join(p.text for p in docx.Document(arqs["docx"]).paragraphs)
    assert cliente.nome in texto and "Equipe Zera Contabilidade" in texto
    with zipfile.ZipFile(arqs["docx"]) as z:          # fontes da Zera embutidas no Word
        assert sum(n.endswith(".odttf") for n in z.namelist()) == 4
    assert arqs["pdf"].stat().st_size > 20_000
