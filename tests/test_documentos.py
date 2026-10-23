"""Importação livre de documentos (Excel, CSV, PDF) — arquivos sintéticos, sem dados de cliente."""
import datetime as dt

import openpyxl
import pandas as pd
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet

from zera_planejamento.base import BaseMensal
from zera_planejamento.dominio.generico import (ErroLeituraDocumento, competencia, destinos, detectar_layout, extrair,
                                                lembrar, ler_tabelas, montar_parcial, sugerir_destino)

P = lambda s: pd.Period(s, "M")  # noqa: E731


def _dest(longo, mapa=None):
    return {r: d for r, (d, _) in destinos(mapa or {}, longo["rotulo"].unique(), "rec_anexo_iii").items()}


def test_competencia_em_varios_formatos():
    assert competencia("jan/26") == P("2026-01")
    assert competencia("Março 2026") == P("2026-03")
    assert competencia("01/2026") == P("2026-01")
    assert competencia("2026-02") == P("2026-02")
    assert competencia("31/12/2025") == P("2025-12")
    assert competencia(dt.datetime(2026, 4, 30)) == P("2026-04")
    assert competencia("Competência: 05/2026") == P("2026-05")
    for nao in ("Receita", "13/2026", 2026, "1.234,56", None):
        assert competencia(nao) is None


def test_excel_dre_acumulada_vira_mensal(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for linha in (["EMPRESA X LTDA"], ["DRE acumulada"], ["Conta", "jan/26", "fev/26", "mar/26"],
                  ["Receita de serviços", 100_000, 210_000, 330_000], ["(-) Custos dos serviços", -20_000, -41_000, -63_000],
                  ["Lucro líquido", 80_000, 169_000, 267_000]):
        ws.append(linha)
    wb.save(tmp_path / "dre.xlsx")
    tab = ler_tabelas(tmp_path / "dre.xlsx")[0]
    lay = detectar_layout(tab)
    assert lay.tipo == "colunas" and len(lay.meses_por_coluna) == 3
    longo = extrair(tab, lay)
    d = _dest(longo)
    assert d["Lucro líquido"] == "ignorar" and d["(-) Custos dos serviços"] == "custos"
    # acumulado → mensal: receita 100 / 110 / 120 mil; custos 20 / 21 / 22 mil (em módulo)
    parc = montar_parcial(longo, d, acumulado=True)
    assert list(parc["rec_anexo_iii"]) == [100_000, 110_000, 120_000]
    assert list(parc["custos"]) == [20_000, 21_000, 22_000]


def test_csv_com_meses_nas_linhas(tmp_path):
    (tmp_path / "fat.csv").write_text("Mês;Faturamento;Devoluções\n01/2026;100.000,00;1.000,00\n"
                                      "02/2026;110.000,00;0\n", encoding="utf-8")
    tab = ler_tabelas(tmp_path / "fat.csv")[0]
    lay = detectar_layout(tab)
    assert lay.tipo == "linhas"
    parc = montar_parcial(extrair(tab, lay), {"Faturamento": "rec_anexo_iii", "Devoluções": "devolucoes"})
    assert parc.loc[P("2026-02"), "rec_anexo_iii"] == 110_000
    assert parc.loc[P("2026-01"), "devolucoes"] == 1_000


def test_pdf_texto_e_pdf_com_grade(tmp_path):
    estilo = getSampleStyleSheet()["Title"]
    SimpleDocTemplate(str(tmp_path / "folha.pdf"), pagesize=A4).build([
        Paragraph("Resumo da folha", estilo),
        Table([["Evento", "01/2026", "02/2026"], ["Salários", "15.000,00", "15.500,00"], ["FGTS", "1.200,00", "1.240,00"]])])
    tab = ler_tabelas(tmp_path / "folha.pdf")[0]
    parc = montar_parcial(extrair(tab, detectar_layout(tab)), {"Salários": "salarios", "FGTS": "fgts"})
    assert parc.loc[P("2026-02"), "salarios"] == 15_500 and parc.loc[P("2026-01"), "fgts"] == 1_200

    grade = Table([["Conta", "Valor"], ["Receita bruta de serviços", "85.000,00"], ["Custos", "(12.000,00)"],
                   ["Lucro líquido", "73.000,00"]])
    grade.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, "black")]))
    SimpleDocTemplate(str(tmp_path / "dre_mes.pdf"), pagesize=A4).build([Paragraph("DRE", estilo), grade])
    tab = ler_tabelas(tmp_path / "dre_mes.pdf")[0]
    lay = detectar_layout(tab)
    assert lay.tipo == "unico" and tab.origem.startswith("Página 1 · tabela")
    longo = extrair(tab, lay, P("2026-03"))
    parc = montar_parcial(longo, _dest(longo))
    assert parc.loc[P("2026-03"), "rec_anexo_iii"] == 85_000 and parc.loc[P("2026-03"), "custos"] == 12_000
    assert "despesas" not in parc.columns          # lucro ignorado; nada inventado


def test_pdf_sem_texto_e_formato_invalido(tmp_path):
    SimpleDocTemplate(str(tmp_path / "imagem.pdf"), pagesize=A4).build([Spacer(1, 10), PageBreak(), Spacer(1, 10)])
    with pytest.raises(ErroLeituraDocumento, match="sem texto"):
        ler_tabelas(tmp_path / "imagem.pdf")
    (tmp_path / "x.docx").write_bytes(b"nada")
    with pytest.raises(ErroLeituraDocumento, match="não suportado"):
        ler_tabelas(tmp_path / "x.docx")


def test_mapa_do_cliente_e_reaplicado_e_vence_a_sugestao():
    assert sugerir_destino("Salários e ordenados") == "salarios"
    assert sugerir_destino("Receita de vendas de mercadorias") == "rec_anexo_i"
    assert sugerir_destino("Faturamento", "rec_fator_r") == "rec_fator_r"
    mapa = lembrar({}, {"Serviços de manutenção": "rec_anexo_iv", "Total geral": "ignorar"})
    escolha = destinos(mapa, ["SERVIÇOS DE MANUTENÇÃO", "Receita de serviços"], "rec_anexo_iii")
    assert escolha["SERVIÇOS DE MANUTENÇÃO"] == ("rec_anexo_iv", "salvo")       # rótulo normalizado (caixa/acento)
    assert escolha["Receita de serviços"] == ("rec_anexo_iii", "sugerido")


def test_gravar_na_base_mescla_so_os_meses_do_documento():
    base = BaseMensal.de_dataframe(pd.DataFrame({"rec_anexo_iii": [90_000.0, 95_000.0]},
                                                index=pd.period_range("2026-01", "2026-02", freq="M")))
    longo = pd.DataFrame({"rotulo": ["Faturamento"], "competencia": [P("2026-02")], "valor": [99_000.0]})
    base.mesclar(montar_parcial(longo, {"Faturamento": "rec_anexo_iii"}), "documento teste.xlsx")
    assert base.dados.loc["2026-01", "rec_anexo_iii"] == 90_000      # janeiro intacto
    assert base.dados.loc["2026-02", "rec_anexo_iii"] == 99_000
    assert base.origem.loc["2026-02", "rec_anexo_iii"] == "documento teste.xlsx"


def test_tela_outros_documentos_do_upload_a_previa(tmp_path, monkeypatch, cliente):
    from pathlib import Path

    from streamlit.testing.v1 import AppTest

    import zera_planejamento.app as app_pkg
    monkeypatch.setenv("ZERA_DADOS", str(tmp_path / "dados"))
    cliente.salvar()
    wb = openpyxl.Workbook()
    for linha in (["Conta", "jan/26", "fev/26"], ["Receita de serviços", 100_000, 110_000],
                  ["Salários", -15_000, -15_000], ["Total", 85_000, 95_000]):
        wb.active.append(linha)
    wb.save(tmp_path / "dre.xlsx")
    at = AppTest.from_file(str(Path(app_pkg.__file__).parent / "main.py"), default_timeout=120)
    at.session_state["cliente_id"] = cliente.id
    at.session_state["_documentos_teste"] = [tmp_path / "dre.xlsx"]
    at.run()
    at.switch_page("paginas/documentos.py").run()
    assert not at.exception, [e.message for e in at.exception]
    assert any("2 meses nas colunas" in i.value for i in at.info)
    assert any(b.label == "Gravar na base do cliente" for b in at.button)
    next(b for b in at.button if b.label == "Gravar na base do cliente").click().run()
    assert not at.exception, [e.message for e in at.exception]
    base = BaseMensal.carregar(cliente.pasta())
    assert base.dados.loc["2026-02", "rec_anexo_iii"] == 110_000 and base.dados.loc["2026-01", "salarios"] == 15_000
    from zera_planejamento.cliente import Cliente
    salvo = Cliente.carregar(cliente.pasta() / "cliente.json")
    assert salvo.mapa_rotulos["total"] == "ignorar"          # escolhas lembradas para o próximo documento


def test_dre_com_rotulos_em_niveis_e_colunas_saldo_soma_total(tmp_path):
    # layout típico de DRE/balancete: rótulo muda de coluna conforme o nível; colunas Saldo | Soma | Total;
    # título com a data → competência sugerida e aviso de acumulado ("exercício")
    wb = openpyxl.Workbook()
    for linha in (["Empresa:", None, "X LTDA", None, None, None, "Folha:", 1],
                  ["DEMONSTRAÇÃO DO RESULTADO DO EXERCÍCIO EM 30/09/2026"],
                  ["Código", "Classificação", "Descrição", None, "Saldo", "Soma", "Total"],
                  [None, None, "RECEITA BRUTA"],
                  [411, "3.1.1.02.000001", None, "SERVIÇOS PRESTADOS", 500_000, 500_000, 500_000],
                  [None, None, "RECEITA LÍQUIDA", None, None, None, 450_000],
                  [480, "3.1.2.03.000008", None, "(-) SIMPLES NACIONAL", -50_000, -50_000, -50_000],
                  [None, None, "LUCRO LÍQUIDO", None, None, None, 300_000],
                  [332, "4.2.2.01.000002", None, "PRÓ-LABORE", -40_000, -40_000, -40_000]):
        wb.active.append(linha)
    wb.save(tmp_path / "dre.xlsx")
    tab = ler_tabelas(tmp_path / "dre.xlsx")[0]
    lay = detectar_layout(tab)
    assert lay.tipo == "unico" and lay.colunas_valor == [4]          # coluna Saldo, não Total
    assert lay.competencia_sugerida == P("2026-09") and lay.talvez_acumulado
    longo = extrair(tab, lay, lay.competencia_sugerida)
    assert list(longo["rotulo"]) == ["SERVIÇOS PRESTADOS", "(-) SIMPLES NACIONAL", "PRÓ-LABORE"]   # sem subtotais


def test_linhas_de_tamanhos_diferentes_nao_quebram(tmp_path):
    (tmp_path / "x.csv").write_text("Relatório\nData;Valor;Obs\n01/2026;10,00\n02/2026;20,00;ok\n", encoding="utf-8")
    tab = ler_tabelas(tmp_path / "x.csv")[0]
    lay = detectar_layout(tab)
    assert montar_parcial(extrair(tab, lay), {"Valor": "despesas"})["despesas"].sum() == 30
