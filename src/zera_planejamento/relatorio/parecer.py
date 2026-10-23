"""
Parecer ao cliente (Word + PDF) com o timbrado da Zera — 1 a 2 páginas em linguagem simples:
o que recomendamos, quanto custa, o que muda e até quando decidir.

O PDF técnico (``pdf.py``) continua sendo o documento interno; este é o que vai para o cliente,
depois da revisão do contador. O layout reproduz o modelo do kit "ZERA - IDENTIDADE VISUAL"
(``zera_doc.js``): A4, margens de 2 cm, faixa do timbrado no topo, rótulos dourados espaçados,
títulos em Cormorant Garamond, texto em Inter 9,5 pt, cartões numerados 01-02-03, tabela sem linhas
verticais e rodapé "ZERA ASSESSORIA EMPRESARIAL E CONTÁBIL" com fio dourado e "página / total".
As fontes vão embutidas no .docx (mesma técnica do ``embutir_fontes.py`` do kit).
"""
from __future__ import annotations

import datetime as dt
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from ..formatos import brl, comp, pct

ESTATICO = Path(__file__).parents[1] / "app" / "static"
FONTES = ESTATICO / "fonts"
TIMBRADO = ESTATICO / "header_a4.png"
OURO, FIO, TINTA, CINZA, FIO_FINO = "A8843F", "C5A15F", "1A1A1A", "5F5F5F", "D9C9A8"
CURTO = {"simples": "Simples Nacional", "simples_hibrido": "Simples Nacional com IBS/CBS por fora",
         "presumido": "Lucro Presumido", "real": "Lucro Real"}
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro",
         "novembro", "dezembro"]

O_QUE_MUDA = {
    "simples": [("Uma guia por mês: ", "o DAS reúne os impostos federais, a contribuição patronal sobre a folha "
                 "(na maioria das atividades) e o ICMS ou ISS."),
                ("Imposto sobre o faturamento: ", "a alíquota sobe conforme a receita dos últimos 12 meses."),
                ("Rotina mais simples: ", "menos obrigações e menos cálculos mês a mês.")],
    "simples_hibrido": [("DAS menor: ", "o IBS e a CBS saem do DAS e passam a ser apurados à parte."),
                        ("Crédito para quem compra de você: ", "clientes empresas aproveitam o crédito integral do "
                         "IBS/CBS, o que pode ajudar nas vendas."),
                        ("Mais controle: ", "a apuração do IBS/CBS por fora exige notas e créditos organizados.")],
    "presumido": [("Imposto de renda sobre um lucro fixo: ", "a lei presume um percentual do faturamento como lucro "
                   "(em geral 32% em serviços e 8% em comércio)."),
                  ("Folha mais cara: ", "a contribuição de 20% ao INSS sobre salários e pró-labore é paga à parte."),
                  ("Bom para margens altas: ", "se o lucro real é maior que o presumido, a diferença não é tributada.")],
    "real": [("Imposto sobre o lucro de verdade: ", "se o lucro cai, o imposto cai junto; prejuízo pode ser "
              "compensado depois (até 30% do lucro de cada período)."),
             ("Créditos nas compras: ", "parte dos impostos pagos em insumos, energia e aluguel volta como crédito."),
             ("Contabilidade em dia é obrigatória: ", "a apuração depende de documentos completos todo mês.")],
}


@dataclass
class Parecer:
    cliente: str
    cnpj: str
    data: dt.date
    rotulo: str
    titulo: str
    contexto: str
    cards: list[tuple[str, str, str]]
    tabela_cab: list[str]
    tabela: list[list[str]]
    destaque: int | None                     # linha da tabela recomendada
    mudancas: list[tuple[str, str]] = field(default_factory=list)
    atencao: list[tuple[str, str]] = field(default_factory=list)
    proximo: list[tuple[str, str]] = field(default_factory=list)
    nota: str = ""
    assinatura: str = "Equipe Zera Contabilidade"


def _data_extenso(d: dt.date) -> str:
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def montar_parecer(pl, cenario: str = "Neutro", hoje: dt.date | None = None, responsavel: str | None = None) -> Parecer:
    """Conteúdo do parecer a partir de um ``Planejamento`` (números do motor, sem recalcular)."""
    from ..painel import montar_painel
    hoje = hoje or dt.date.today()
    p = montar_painel(pl, cenario, "projecao")
    ind, c = p.indicadores, pl.cliente
    mh, mp = pl.meses_historico, pl.meses_projetados
    nome = c.nome
    if "recomendado" not in ind:
        raise ValueError("Sem regime calculável para este cliente: complete a base antes de gerar o parecer.")
    rec, atual = ind["recomendado"], ind.get("atual")
    mudar = ind["mudar"]

    contexto = (f"Comparamos quanto a {nome} pagaria de impostos em cada regime de tributação, com base nos números "
                f"de {comp(mh[0])} a {comp(mh[-1])} e na projeção de {comp(mp[0])} a {comp(mp[-1])}. Os valores "
                "consideram as regras de 2026 e o início da Reforma Tributária em 2027, e estão anualizados para "
                "facilitar a comparação.")
    if mudar and "economia_ano" in ind:
        card1 = ("01", "Recomendação", f"Mudar para o {CURTO[rec]}.")
        card2 = ("02", "Economia estimada", f"{brl(ind['economia_ano'], 0)} por ano em relação ao {CURTO.get(atual, '')}.")
    else:
        card1 = ("01", "Recomendação", f"Manter o {CURTO[rec]}.")
        extra = (f" A diferença para o {ind['segundo'].replace(' (IBS/CBS por fora)', ' com IBS/CBS por fora')} é de "
                 f"{brl(ind['vantagem_segundo'] * 12 / ind['n_meses'], 0)} por ano.") if "segundo" in ind else ""
        card2 = ("02", "Carga estimada", f"{pct(ind['carga_recomendado'], 1)} do faturamento.{extra}")
    if ind.get("prazo_acao") == "DECIDIR":
        card3 = ("03", "Prazo", f"Decidir até {ind['prazo_data']:%d/%m/%Y}. A opção vale para o ano inteiro.")
    else:
        card3 = ("03", "Próxima revisão", "Reavaliar quando a alíquota da CBS de 2027 for publicada e no fechamento "
                 "do ano.")

    linhas, destaque = [], None
    for i, r in enumerate(p.ranking.itertuples()):
        situacao = "Recomendado" if r.melhor else ("Atual" if r.atual else "")
        if r.melhor and r.atual:
            situacao = "Atual e recomendado"
        if r.melhor:
            destaque = i
        linhas.append([CURTO[r.regime], brl(r.tributos_ano, 0), pct(r.carga, 1), situacao])

    atencao = []
    if "vantagem_segundo_pct" in ind and ind["vantagem_segundo_pct"] < 0.03:
        atencao.append(("Diferença pequena: ", "a vantagem sobre o segundo regime é menor que 3%; pequenas mudanças no "
                        "faturamento ou nos custos podem inverter a ordem."))
    if ind.get("socios_muda"):
        atencao.append(("Retirada dos sócios: ", f"somando o imposto sobre lucros distribuídos (Lei 15.270/2025), o "
                        f"menor custo total passa a ser o {ind['socios_melhor']}. Vale conversarmos sobre a forma de "
                        "remunerar os sócios."))
    elif pl.socios is not None and not pl.socios.vazia and pl.socios.detalhe["dividendos_mes"].max() > 50_000:
        atencao.append(("Lucros acima de R$ 50 mil por mês: ", "desde 2026 há retenção de 10% na fonte por sócio "
                        "(Lei 15.270/2025). Podemos planejar a distribuição."))
    if "fator_r" in ind and rec in ("simples", "simples_hibrido"):
        fr = ind["fator_r"]
        if fr < 0.28:
            atencao.append(("Fator R: ", f"a folha representa {pct(fr, 1)} do faturamento. Acima de 28%, os serviços "
                            "pagam menos imposto no Simples; avaliamos o pró-labore ideal com vocês."))
        else:
            atencao.append(("Fator R: ", f"a folha está em {pct(fr, 1)} do faturamento, acima dos 28% que garantem a "
                            "alíquota menor. Reduções de folha ou de pró-labore precisam ser avaliadas antes."))
    atencao.append(("Reforma Tributária: ", "a alíquota da CBS de 2027 ainda será fixada pelo Senado; os valores de "
                    "2027 em diante são estimativas e serão atualizados."))
    if ind["alertas"]["CRÍTICO"]:
        atencao.append(("Dados em conferência: ", "alguns números ainda estão sendo conferidos pela nossa equipe."))

    if mudar and ind.get("prazo_acao") == "DECIDIR":
        proximo = [("Confirme a decisão ", f"até {ind['prazo_data'] - dt.timedelta(days=7):%d/%m/%Y}, para termos "
                    "tempo de fazer a opção dentro do prazo."),
                   ("Nós cuidamos da mudança: ", "opção do regime, ajustes no sistema e orientação sobre as novas guias.")]
    elif mudar:
        proximo = [("Vamos conversar ", "sobre a mudança e o melhor momento para fazê-la."),
                   ("Nós cuidamos da mudança: ", "opção do regime, ajustes no sistema e orientação sobre as novas guias.")]
    else:
        proximo = [("Nenhuma ação agora: ", "seguimos no regime atual."),
                   ("Acompanhamento: ", "refazemos a análise se o faturamento, a folha ou a margem mudarem bastante.")]

    return Parecer(
        cliente=nome, cnpj=c.cnpj, data=hoje,
        rotulo=f"Planejamento tributário · {MESES[hoje.month - 1]} de {hoje.year}",
        titulo="Qual regime de tributação custa menos", contexto=contexto, cards=[card1, card2, card3],
        tabela_cab=["Regime", "Impostos por ano", "Sobre a receita", ""], tabela=linhas, destaque=destaque,
        mudancas=O_QUE_MUDA.get(rec, []), atencao=atencao, proximo=proximo,
        nota=("Estimativa feita com os dados informados e a legislação vigente nesta data, no cenário "
              f"{cenario.lower()}. Não substitui a apuração oficial. Impostos incluem DAS, PIS/Cofins ou CBS/IBS, "
              "IRPJ, CSLL, ICMS, ISS e INSS patronal, conforme o regime."),
        assinatura=responsavel or "Equipe Zera Contabilidade")


# =====================================================================================================
# Word (.docx)
# =====================================================================================================
def _docx():
    import docx  # noqa: F401 — importação tardia: só quem gera o parecer precisa do python-docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Twips
    return docx, WD_TABLE_ALIGNMENT, WD_TAB_ALIGNMENT, OxmlElement, qn, Pt, RGBColor, Twips


def exportar_parecer_docx(par: Parecer, caminho: Path) -> Path:
    docx, WD_TABLE_ALIGNMENT, WD_TAB_ALIGNMENT, OxmlElement, qn, Pt, RGBColor, Twips = _docx()
    LARG = 9638                                   # largura útil em twips (A4 − 2 × 2 cm)
    doc = docx.Document()

    def fonte(run, nome="Inter", tam=9.5, cor=TINTA, negrito=False, espaco=None):
        run.font.name, run.font.size, run.font.bold = nome, Pt(tam), negrito
        run.font.color.rgb = RGBColor.from_string(cor)
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(k), nome)
        if espaco:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(espaco))
            rpr.append(sp)
        return run

    def paragrafo(antes=0, depois=100, linha=276, recuo=None, manter=False):
        p = doc.add_paragraph()
        f = p.paragraph_format
        f.space_before, f.space_after, f.keep_with_next = Twips(antes), Twips(depois), manter
        f.line_spacing = linha / 240
        if recuo is not None:
            f.left_indent, f.first_line_indent = Twips(recuo), Twips(-recuo)
        return p

    def borda(el, lados: dict):
        tag = "w:tcBorders" if el.tag == qn("w:tcPr") else "w:pBdr"
        b = OxmlElement(tag)
        for lado, (tam, cor) in lados.items():
            e = OxmlElement(f"w:{lado}")
            e.set(qn("w:val"), "single" if tam else "nil")
            if tam:
                e.set(qn("w:sz"), str(tam))
                e.set(qn("w:color"), cor)
                e.set(qn("w:space"), "6" if tag == "w:pBdr" else "0")
            b.append(e)
        el.append(b)

    def tabela_base(n_lin, larguras):
        t = doc.add_table(rows=n_lin, cols=len(larguras))
        t.alignment, t.autofit = WD_TABLE_ALIGNMENT.LEFT, False
        tblpr = t._tbl.tblPr
        for e in list(tblpr):
            if e.tag in (qn("w:tblBorders"), qn("w:tblStyle")):
                tblpr.remove(e)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblpr.append(lay)
        for linha in t.rows:
            for j, cel in enumerate(linha.cells):
                cel.width = Twips(larguras[j])
        return t

    def margens(cel, topo=90, base=90, esq=0, dir_=100):
        tcpr = cel._tc.get_or_add_tcPr()
        m = OxmlElement("w:tcMar")
        for lado, v in (("top", topo), ("bottom", base), ("left", esq), ("right", dir_)):
            e = OxmlElement(f"w:{lado}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            m.append(e)
        tcpr.append(m)
        return tcpr

    # página, cabeçalho (faixa do timbrado) e rodapé
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.top_margin, sec.bottom_margin = Twips(2150), Twips(1250)
    sec.left_margin = sec.right_margin = Twips(1134)
    sec.header_distance, sec.footer_distance = Twips(0), Twips(500)
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size = "Inter", Pt(9.5)
    estilo.element.get_or_add_rPr()
    if TIMBRADO.exists():
        hp = sec.header.paragraphs[0]
        hp.paragraph_format.left_indent = Twips(-1134)
        hp.paragraph_format.space_after = Twips(0)
        hp.add_run().add_picture(str(TIMBRADO), width=Twips(11906))
    fp = sec.footer.paragraphs[0]
    borda(fp._p.get_or_add_pPr(), {"top": (6, FIO)})
    for parada in doc.styles["Footer"].paragraph_format.tab_stops:   # o estilo padrão centraliza o 1º tab
        fp.paragraph_format.tab_stops.add_tab_stop(parada.position, WD_TAB_ALIGNMENT.CLEAR)
    fp.paragraph_format.tab_stops.add_tab_stop(Twips(LARG), WD_TAB_ALIGNMENT.RIGHT)
    fonte(fp.add_run("ZERA ASSESSORIA EMPRESARIAL E CONTÁBIL"), tam=6.5, cor="8A8A8A", espaco=30)
    fp.add_run("\t")
    for instr in ("PAGE", None, "NUMPAGES"):
        if instr is None:
            fonte(fp.add_run(" / "), tam=7, cor=OURO)
            continue
        campo = OxmlElement("w:fldSimple")
        campo.set(qn("w:instr"), instr)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        campo.append(r)
        fp._p.append(campo)

    def rotulo(txt, antes=220):
        fonte(paragrafo(antes, 100, manter=True).add_run(txt.upper()), tam=7.5, cor=OURO, espaco=40)

    def item(neg, txt):
        p = paragrafo(0, 60, 252, recuo=300)
        fonte(p.add_run("•\t"), "Arial", 10, FIO)
        p.paragraph_format.tab_stops.add_tab_stop(Twips(300))
        fonte(p.add_run(neg), tam=9, negrito=True)
        fonte(p.add_run(txt), tam=9, cor=CINZA)

    # conteúdo
    rotulo(par.rotulo, antes=0)
    fonte(paragrafo(0, 120, 300, manter=True).add_run(par.titulo), "Cormorant Garamond", 22)
    p = paragrafo(0, 140, 288)
    fonte(p.add_run("Para: "), cor=CINZA)
    fonte(p.add_run(par.cliente), negrito=True)
    if par.cnpj:
        fonte(p.add_run(f"  ·  CNPJ {par.cnpj}"), cor=CINZA)
    fonte(p.add_run(f"  ·  São Paulo, {_data_extenso(par.data)}"), cor=CINZA)
    fonte(paragrafo(0, 140, 276).add_run(par.contexto))

    larg_card = [LARG // 3] * 3
    larg_card[-1] += LARG - sum(larg_card)
    tc = tabela_base(1, larg_card)
    for j, (num, tit, txt) in enumerate(par.cards):
        cel = tc.rows[0].cells[j]
        tcpr = margens(cel, 120, 60, 160 if j else 0, 160)
        borda(tcpr, {"top": (6, FIO), "left": (0, ""), "bottom": (0, ""), "right": (0, "")})
        p0 = cel.paragraphs[0]
        p0.paragraph_format.space_after = Twips(20)
        fonte(p0.add_run(num), "Cormorant Garamond", 13, OURO)
        p1 = cel.add_paragraph()
        p1.paragraph_format.space_after, p1.paragraph_format.line_spacing = Twips(40), 1.0
        fonte(p1.add_run(tit), "Cormorant Garamond", 14)
        p2 = cel.add_paragraph()
        p2.paragraph_format.space_after, p2.paragraph_format.line_spacing = Twips(0), 1.1
        fonte(p2.add_run(txt), tam=8.5, cor=CINZA)

    rotulo("Comparação dos regimes")
    larg = [3300, 1900, 2050, LARG - 7250]
    tt = tabela_base(len(par.tabela) + 1, larg)
    for i, linha in enumerate([par.tabela_cab] + par.tabela):
        for j, valor in enumerate(linha):
            cel = tt.rows[i].cells[j]
            tcpr = margens(cel, 60, 60, 100 if j else 0, 100)
            borda(tcpr, {"top": (0, ""), "left": (0, ""), "right": (0, ""),
                         "bottom": (6, FIO) if i == 0 else (4, FIO_FINO)})
            pc = cel.paragraphs[0]
            pc.paragraph_format.space_after = Twips(0)
            if i == 0:
                fonte(pc.add_run(valor.upper()), tam=7, cor=OURO, espaco=30)
            else:
                destaque = par.destaque is not None and i - 1 == par.destaque
                fonte(pc.add_run(valor), tam=8.5, negrito=destaque, cor=OURO if destaque and j == 3 else TINTA)

    if par.mudancas:
        rotulo("O que muda na prática")
        for neg, txt in par.mudancas:
            item(neg, txt)
    rotulo("Pontos de atenção")
    for neg, txt in par.atencao:
        item(neg, txt)
    rotulo("Próximo passo")
    for neg, txt in par.proximo:
        item(neg, txt)
    fonte(paragrafo(200, 0).add_run(par.assinatura), "Cormorant Garamond", 15, OURO)
    fonte(paragrafo(160, 0, 252).add_run(par.nota), tam=7.5, cor=CINZA)

    caminho = Path(caminho)
    doc.save(caminho)
    _embutir_fontes(caminho)
    return caminho


def _embutir_fontes(caminho: Path) -> None:
    """Embute Cormorant Garamond e Inter no .docx (fontes ofuscadas, ECMA-376 Parte 2), como o kit da Zera."""
    familias = {"Cormorant Garamond": ("CormorantGaramond-Regular.ttf", "CormorantGaramond-Bold.ttf"),
                "Inter": ("Inter-Regular.ttf", "Inter-Bold.ttf")}
    if not all((FONTES / f).exists() for par_ in familias.values() for f in par_):
        return
    with zipfile.ZipFile(caminho) as z:
        arquivos = {n: z.read(n) for n in z.namelist()}

    def ofuscar(dados: bytes, guid: str) -> bytes:
        chave = bytes.fromhex(guid.strip("{}").replace("-", ""))[::-1]
        b = bytearray(dados)
        for i in range(32):
            b[i] ^= chave[i % 16]
        return bytes(b)

    rels, entradas, n = [], [], 0
    for familia, (regular, negrito) in familias.items():
        emb = ""
        for estilo, arq in (("Regular", regular), ("Bold", negrito)):
            n += 1
            guid = "{" + str(uuid.uuid4()).upper() + "}"
            arquivos[f"word/fonts/font{n}.odttf"] = ofuscar((FONTES / arq).read_bytes(), guid)
            rels.append(f'<Relationship Id="rIdF{n}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                        f'relationships/font" Target="fonts/font{n}.odttf"/>')
            emb += f'<w:embed{estilo} r:id="rIdF{n}" w:fontKey="{guid}"/>'
        tipo = "roman" if "Cormorant" in familia else "swiss"
        entradas.append(f'<w:font w:name="{familia}"><w:charset w:val="00"/><w:family w:val="{tipo}"/>'
                        f'<w:pitch w:val="variable"/>{emb}</w:font>')
    ft = arquivos["word/fontTable.xml"].decode("utf-8")
    for familia in familias:
        ft = re.sub(r'<w:font w:name="%s">.*?</w:font>' % re.escape(familia), "", ft, flags=re.S)
    if 'xmlns:r="' not in ft.split(">", 2)[1]:
        ft = ft.replace("<w:fonts ", '<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
                        'relationships" ', 1)
    arquivos["word/fontTable.xml"] = ft.replace("</w:fonts>", "".join(entradas) + "</w:fonts>").encode("utf-8")
    arquivos["word/_rels/fontTable.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.'
        'org/package/2006/relationships">' + "".join(rels) + "</Relationships>").encode("utf-8")
    ct = arquivos["[Content_Types].xml"].decode("utf-8")
    if "odttf" not in ct:
        ct = re.sub(r"(<Types[^>]*>)", r'\1<Default Extension="odttf" ContentType="application/vnd.openxmlformats-'
                    r'officedocument.obfuscatedFont"/>', ct, count=1)
    arquivos["[Content_Types].xml"] = ct.encode("utf-8")
    st = arquivos["word/settings.xml"].decode("utf-8")
    if "embedTrueTypeFonts" not in st:
        st = re.sub(r"(<w:settings[^>]*>)", r"\1<w:embedTrueTypeFonts/>", st, count=1)
    arquivos["word/settings.xml"] = st.encode("utf-8")
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", arquivos.pop("[Content_Types].xml"))
        for nome, dados in arquivos.items():
            z.writestr(nome, dados)


# =====================================================================================================
# PDF
# =====================================================================================================
def exportar_parecer_pdf(par: Parecer, caminho: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    serif, sans, sans_b = "Times-Roman", "Helvetica", "Helvetica-Bold"
    if (FONTES / "Inter-Regular.ttf").exists():
        for nome, arq in (("ZeraSerif", "CormorantGaramond-Regular.ttf"), ("ZeraSans", "Inter-Regular.ttf"),
                          ("ZeraSans-Bold", "Inter-Bold.ttf")):
            if nome not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(nome, str(FONTES / arq)))
        serif, sans, sans_b = "ZeraSerif", "ZeraSans", "ZeraSans-Bold"
        pdfmetrics.registerFontFamily("ZeraSans", normal=sans, bold=sans_b, italic=sans, boldItalic=sans_b)
    ouro, fio, tinta, cinza = (colors.HexColor("#" + x) for x in (OURO, FIO, TINTA, CINZA))
    est = {
        "rot": ParagraphStyle("rot", fontName=sans_b, fontSize=7.5, leading=10, textColor=ouro, spaceBefore=14,
                              spaceAfter=5),
        "tit": ParagraphStyle("tit", fontName=serif, fontSize=24, leading=27, textColor=tinta, spaceAfter=6),
        "txt": ParagraphStyle("txt", fontName=sans, fontSize=9.5, leading=13.7, textColor=tinta, spaceAfter=7),
        "num": ParagraphStyle("num", fontName=serif, fontSize=14, leading=16, textColor=ouro),
        "ctit": ParagraphStyle("ctit", fontName=serif, fontSize=15, leading=17, textColor=tinta, spaceAfter=2),
        "ctxt": ParagraphStyle("ctxt", fontName=sans, fontSize=8.5, leading=11.5, textColor=cinza),
        "cab": ParagraphStyle("cab", fontName=sans_b, fontSize=7, leading=9, textColor=ouro),
        "cel": ParagraphStyle("cel", fontName=sans, fontSize=8.5, leading=11, textColor=tinta),
        "celb": ParagraphStyle("celb", fontName=sans_b, fontSize=8.5, leading=11, textColor=tinta),
        "item": ParagraphStyle("item", fontName=sans, fontSize=9, leading=12.5, textColor=cinza, leftIndent=12,
                               spaceAfter=4, bulletColor=fio, bulletFontName="Helvetica", alignment=TA_LEFT),
        "ass": ParagraphStyle("ass", fontName=serif, fontSize=17, leading=20, textColor=ouro, spaceBefore=12),
        "nota": ParagraphStyle("nota", fontName=sans, fontSize=7.5, leading=10, textColor=cinza, spaceBefore=10),
    }

    class Folha(rl_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            w, h = A4
            for estado in self._paginas:
                self.__dict__.update(estado)
                if TIMBRADO.exists():
                    self.drawImage(str(TIMBRADO), 0, h - w * 390 / 2480, width=w, height=w * 390 / 2480, mask="auto")
                self.setStrokeColor(fio)
                self.setLineWidth(0.6)
                self.line(2 * cm, 1.45 * cm, w - 2 * cm, 1.45 * cm)
                self.setFont(sans, 6.5)
                self.setFillColor(colors.HexColor("#8A8A8A"))
                self.drawString(2 * cm, 1.05 * cm, "ZERA ASSESSORIA EMPRESARIAL E CONTÁBIL")
                self.setFillColor(ouro)
                self.setFont(sans, 7)
                self.drawRightString(w - 2 * cm, 1.05 * cm, f"{self._pageNumber} / {total}")
                super().showPage()
            super().save()

    def esc(t: str) -> str:
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    largura = A4[0] - 4 * cm
    doc = SimpleDocTemplate(str(caminho), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=3.8 * cm,
                            bottomMargin=2.2 * cm, title=f"Parecer — {par.cliente}", author=par.assinatura)
    h = [Paragraph(esc(par.rotulo.upper()), est["rot"]), Paragraph(esc(par.titulo), est["tit"]),
         Paragraph(f'<font color="#{CINZA}">Para:</font> <b>{esc(par.cliente)}</b>'
                   + (f'<font color="#{CINZA}">  ·  CNPJ {esc(par.cnpj)}</font>' if par.cnpj else "")
                   + f'<font color="#{CINZA}">  ·  São Paulo, {_data_extenso(par.data)}</font>', est["txt"]),
         Paragraph(esc(par.contexto), est["txt"]), Spacer(1, 4)]
    cards = Table([[[Paragraph(n, est["num"]), Paragraph(esc(t), est["ctit"]), Paragraph(esc(x), est["ctxt"])]
                    for n, t, x in par.cards]], colWidths=[largura / 3] * 3)
    cards.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.8, fio), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("LEFTPADDING", (0, 0), (0, 0), 0), ("LEFTPADDING", (1, 0), (-1, 0), 8),
                               ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 6)]))
    h.append(cards)

    linhas = [[Paragraph(c.upper(), est["cab"]) for c in par.tabela_cab]]
    for i, ln in enumerate(par.tabela):
        st_ = est["celb"] if i == par.destaque else est["cel"]
        linhas.append([Paragraph(esc(v), st_) for v in ln])
    tab = Table(linhas, colWidths=[largura * 0.37, largura * 0.23, largura * 0.21, largura * 0.19], repeatRows=1,
                hAlign="LEFT")
    tab.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.8, fio),
                             ("LINEBELOW", (0, 1), (-1, -1), 0.4, colors.HexColor("#" + FIO_FINO)),
                             ("LEFTPADDING", (0, 0), (0, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 4),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    h += [Paragraph("COMPARAÇÃO DOS REGIMES", est["rot"]), tab]

    def bloco(titulo, itens):
        if not itens:
            return
        h.append(KeepTogether([Paragraph(titulo, est["rot"])] + [
            Paragraph(f'<font color="#{TINTA}"><b>{esc(n)}</b></font>{esc(t)}', est["item"], bulletText="•")
            for n, t in itens]))

    bloco("O QUE MUDA NA PRÁTICA", par.mudancas)
    bloco("PONTOS DE ATENÇÃO", par.atencao)
    bloco("PRÓXIMO PASSO", par.proximo)
    h += [Paragraph(esc(par.assinatura), est["ass"]), Paragraph(esc(par.nota), est["nota"])]
    doc.build(h, canvasmaker=Folha)
    return Path(caminho)


def gerar_parecer(pl, pasta: Path, cenario: str = "Neutro", responsavel: str | None = None) -> dict[str, Path]:
    """Gera ``Parecer_<cliente>.docx`` e ``.pdf`` na pasta; devolve os caminhos."""
    from ..formatos import slug
    par = montar_parecer(pl, cenario, responsavel=responsavel)
    pasta.mkdir(parents=True, exist_ok=True)
    nome = f"Parecer_{slug(pl.cliente.nome)[:40]}_{par.data:%Y%m%d}"
    return {"docx": exportar_parecer_docx(par, pasta / f"{nome}.docx"),
            "pdf": exportar_parecer_pdf(par, pasta / f"{nome}.pdf")}


__all__ = ["Parecer", "montar_parecer", "exportar_parecer_docx", "exportar_parecer_pdf", "gerar_parecer"]
